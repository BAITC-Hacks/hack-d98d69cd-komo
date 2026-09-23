import io
import json
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.core.errors import AppError
from app.features.data_import.formats import TEMPLATE_DIR, decode_file, read_excel, read_json
from app.features.data_import.schemas import EmployeeBundle
from app.features.data_import.service import parse_history
from conftest import login


def template_book():
    return load_workbook(TEMPLATE_DIR / 'career-quest-template.xlsx')


def dump(workbook):
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def test_excel_json_templates_are_equivalent():
    excel = read_excel((TEMPLATE_DIR / 'career-quest-template.xlsx').read_bytes())
    js = read_json((TEMPLATE_DIR / 'career-quest-template.json').read_bytes(), 'employees')
    assert EmployeeBundle.model_validate_json(excel[0]) == EmployeeBundle.model_validate_json(js[0])
    assert parse_history(excel[1]) == parse_history(js[1])


def test_excel_reordered_columns_and_empty_rows():
    book = template_book()
    sheet = book['employees']
    for row in sheet.iter_rows():
        first, second = row[0].value, row[1].value
        row[0].value, row[1].value = second, first
    sheet.insert_rows(2)
    employees, history = read_excel(dump(book))
    assert len(json.loads(employees)['employees']) == 1
    assert len(parse_history(history)) == 1


@pytest.mark.parametrize('column,value,code', [('level', 6, 'invalid_excel_skill'), ('level', '=2+2', 'invalid_excel_cell')])
def test_invalid_excel_skill_has_row(column, value, code):
    book = template_book()
    book['employee_skills']['C2'] = value
    with pytest.raises(AppError) as error:
        read_excel(dump(book))
    assert error.value.code == code
    assert 'строка 2' in error.value.message


def test_excel_unknown_header_and_corrupt_file():
    book = template_book(); book['employees']['A1'] = 'wrong_id'
    with pytest.raises(AppError) as error:
        read_excel(dump(book))
    assert error.value.code == 'invalid_excel_header'
    with pytest.raises(AppError) as error:
        decode_file(b'not a workbook', 'data.xlsx', 'employees')
    assert error.value.code == 'invalid_excel'


def test_history_json_and_flexible_csv():
    data = json.loads((TEMPLATE_DIR / 'career-quest-template.json').read_bytes())
    _, history = read_json(json.dumps(data['history']).encode(), 'history')
    rows = parse_history(history)
    headers = list(reversed(rows[0]))
    # Excel's regional CSV separator and reordered columns.
    csv = ';'.join(headers) + '\n' + ';'.join(str(rows[0][k]) if rows[0][k] is not None else '' for k in headers)
    assert parse_history(csv.encode()) == rows
    employees, _ = read_json(json.dumps(data['employees']).encode(), 'employees')
    assert EmployeeBundle.model_validate_json(employees).employees[0].employee_id == 'TEMPLATE_EMP_001'


async def test_template_rights_excel_import_repeat_and_atomic_failure(client):
    await login(client)
    assert (await client.get('/api/v1/hr/import/template')).status_code == 403
    await login(client, 'hr')
    response = await client.get('/api/v1/hr/import/template?format=xlsx')
    assert response.status_code == 200
    assert 'career-quest-template.xlsx' in response.headers['content-disposition']
    book = load_workbook(io.BytesIO(response.content))
    eid = 'QA_R2_XLSX_' + uuid4().hex[:10]
    for sheet in [book['employees'], book['employee_skills'], book['history']]:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value == 'TEMPLATE_EMP_001': cell.value = eid
                if cell.value == 'TEMPLATE_REC_001': cell.value = eid + '_REC'
    raw = dump(book)
    files = {'employees': ('import.xlsx', raw)}
    first = await client.post('/api/v1/hr/import', files=files)
    assert first.status_code == 200, first.text
    assert first.json()['employees_created'] == first.json()['history_created'] == 1
    repeated = await client.post('/api/v1/hr/import', files=files)
    assert repeated.status_code == 200
    assert repeated.json()['employees_created'] == repeated.json()['history_created'] == 0
    book = load_workbook(io.BytesIO(raw))
    badid = eid + '_BAD'
    for sheet in [book['employees'], book['employee_skills'], book['history']]:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value == eid: cell.value = badid
    book['history']['A2'] = badid + '_REC'
    book['history']['C2'] = 'EV_UNKNOWN'
    failed = await client.post('/api/v1/hr/import', files={'employees': ('bad.xlsx', dump(book))})
    assert failed.status_code == 422
    assert (await client.get(f'/api/v1/employees/{badid}/development')).status_code == 404


async def test_combined_json_and_separate_history(client):
    await login(client, 'hr')
    response = await client.get('/api/v1/hr/import/template?format=json')
    assert response.status_code == 200
    data = response.json(); eid = 'QA_R2_JSON_' + uuid4().hex[:10]
    data['employees'][0]['employee_id'] = eid
    data['history'][0].update(employee_id=eid, record_id=eid + '_REC')
    result = await client.post('/api/v1/hr/import', files={'employees': ('all.json', json.dumps(data))})
    assert result.status_code == 200, result.text
    assert result.json()['history_created'] == result.json()['employees_created'] == 1
    repeat = await client.post('/api/v1/hr/import', files={'history': ('history.json', json.dumps(data['history']))})
    assert repeat.status_code == 200 and repeat.json()['history_created'] == 0
    double = await client.post('/api/v1/hr/import', files={'employees': ('all.json', json.dumps(data)), 'history': ('history.json', json.dumps(data['history']))})
    assert double.status_code == 422 and double.json()['code'] == 'duplicate_source'

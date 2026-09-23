"""File adapters: every format reaches the same atomic domain importer."""
import csv
import io
import json
from datetime import date, datetime
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from pydantic import ValidationError

from app.core.errors import AppError
from app.features.data_import.service import HISTORY_COLUMNS, validation_details
from app.features.employees.schemas import EmployeeProfile
from app.features.development.schemas import HistoryInput

META = {'dataset': 'Career Quest', 'version': '1.0', 'as_of_date': '2026-10-01'}
EMPLOYEE_COLUMNS = ['employee_id', 'full_name', 'department', 'role', 'grade', 'manager_id', 'hire_date', 'tenure_months', 'work_format', 'preferred_language', 'target_role', 'target_grade', 'last_review_date']
SKILL_COLUMNS = ['employee_id', 'skill_id', 'level']
TEMPLATE_DIR = Path(__file__).with_name('templates')


def encode_history(rows):
    if not isinstance(rows, list):
        raise AppError('invalid_history', 'history должен быть массивом записей', 422)
    if len(rows) > 20000:
        raise AppError('import_too_large', 'Не более 20 000 записей за импорт', 413)
    stream = io.StringIO()
    writer = csv.DictWriter(stream, fieldnames=HISTORY_COLUMNS)
    writer.writeheader()
    for i, row in enumerate(rows, 1):
        try:
            normalized = HistoryInput.model_validate(row).model_dump(mode='json')
        except ValidationError as error:
            raise AppError('invalid_history', f'История: запись {i}', 422, validation_details(error))
        writer.writerow(normalized)
    return stream.getvalue().encode() if rows else None


def read_json(raw, kind):
    try:
        value = json.loads(raw.decode('utf-8-sig'))
    except (ValueError, UnicodeDecodeError):
        raise AppError('invalid_json', 'Не удалось прочитать JSON. Проверьте синтаксис и кодировку UTF-8.', 422)
    if isinstance(value, list):
        value = {kind: value}
    if not isinstance(value, dict) or set(value) - {'meta', 'employees', 'history'}:
        raise AppError('invalid_json', 'Ожидается массив или объект с employees, history и необязательным meta.', 422)
    employee_bytes = None
    if 'employees' in value:
        if not isinstance(value['employees'], list):
            raise AppError('invalid_employees', 'employees должен быть массивом профилей', 422)
        if value['employees']:
            employee_bytes = json.dumps({'meta': value.get('meta', META), 'employees': value['employees']}).encode()
    history_bytes = encode_history(value['history']) if 'history' in value else None
    return employee_bytes, history_bytes


def table(sheet, columns, limit):
    """Headers can be reordered; blank rows/columns are harmless, typos are not."""
    if (sheet.max_row or 0) > 200001 or (sheet.max_column or 0) > 100:
        raise AppError('import_too_large', f'Лист {sheet.title}: слишком большая область листа', 413)
    iterator = sheet.iter_rows()
    cells = next(iterator, ())
    header = [str(c.value).strip() if c.value is not None else '' for c in cells]
    names = [h for h in header if h]
    if len(names) != len(set(names)) or set(names) != set(columns):
        raise AppError('invalid_excel_header', f'Лист {sheet.title}: проверьте названия колонок в первой строке', 422,
                       {'required_columns': columns, 'missing': sorted(set(columns) - set(names)), 'unknown': sorted(set(names) - set(columns))})
    count = 0
    for number, cells in enumerate(iterator, 2):
        if number > 200001 or len(cells) > 100:
            raise AppError('import_too_large', f'Лист {sheet.title}: слишком большая область листа', 413)
        if all(c.value is None or c.value == '' for c in cells):
            continue
        count += 1
        if count > limit:
            raise AppError('import_too_large', f'Лист {sheet.title}: не более {limit} непустых строк', 413)
        if any(c.value is not None for c in cells[len(header):]):
            raise AppError('invalid_excel_header', f'Лист {sheet.title}, строка {number}: значение в колонке без заголовка', 422)
        row = {key: None for key in names}
        for key, cell in zip(header, cells):
            if not key:
                if cell.value is not None:
                    raise AppError('invalid_excel_header', f'Лист {sheet.title}, строка {number}: значение в колонке без заголовка', 422)
                continue
            if cell.data_type in ('f', 'e'):
                raise AppError('invalid_excel_cell', f'Лист {sheet.title}, строка {number}, поле {key}: нужны значения, не формулы или ошибки Excel', 422)
            value = cell.value
            if isinstance(value, datetime):
                value = value.date().isoformat()
            elif isinstance(value, date):
                value = value.isoformat()
            elif isinstance(value, str):
                value = value.strip() or None
            row[key] = value
        yield number, row


def read_excel(raw):
    try:
        with ZipFile(io.BytesIO(raw)) as archive:
            if sum(f.file_size for f in archive.infolist()) > 32 * 1024 * 1024:
                raise AppError('file_too_large', 'Распакованный Excel не должен превышать 32 МБ', 413)
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except AppError:
        raise
    except Exception:
        raise AppError('invalid_excel', 'Не удалось прочитать Excel. Сохраните книгу в формате .xlsx без пароля.', 422)
    try:
        if not {'employees', 'history'}.intersection(workbook.sheetnames):
            raise AppError('invalid_excel_sheets', 'В Excel нужен лист employees и/или history. Используйте скачанный шаблон.', 422)
        profiles, skill_map = [], {}
        if 'employee_skills' in workbook:
            for number, row in table(workbook['employee_skills'], SKILL_COLUMNS, 120000):
                eid, skill, level = row['employee_id'], row['skill_id'], row['level']
                if not isinstance(eid, str) or not isinstance(skill, str) or type(level) is not int or not 0 <= level <= 5:
                    raise AppError('invalid_excel_skill', f'Лист employee_skills, строка {number}: employee_id и skill_id — текст, level — целое число от 0 до 5', 422)
                skills = skill_map.setdefault(eid, {})
                if skill in skills:
                    raise AppError('duplicate_skill', f'Лист employee_skills, строка {number}: навык повторяется у сотрудника {eid}', 422)
                skills[skill] = level
        if 'employees' in workbook:
            for number, row in table(workbook['employees'], EMPLOYEE_COLUMNS, 2000):
                role, grade = row.pop('target_role'), row.pop('target_grade')
                row['career_goal'] = {'target_role': role, 'target_grade': grade} if role or grade else None
                row['skills'] = skill_map.pop(row.get('employee_id'), {})
                try:
                    profiles.append(EmployeeProfile.model_validate(row).model_dump(mode='json'))
                except ValidationError as error:
                    raise AppError('invalid_employees', f'Лист employees, строка {number}', 422, validation_details(error))
        if skill_map:
            raise AppError('unknown_skill_employee', 'На листе employee_skills есть ID, отсутствующий на листе employees', 422, {'employee_ids': list(skill_map)[:10]})
        history = []
        if 'history' in workbook:
            for number, row in table(workbook['history'], HISTORY_COLUMNS, 20000):
                try:
                    history.append(HistoryInput.model_validate(row).model_dump(mode='json'))
                except ValidationError as error:
                    raise AppError('invalid_history', f'Лист history, строка {number}', 422, validation_details(error))
        return (json.dumps({'meta': META, 'employees': profiles}).encode() if profiles else None, encode_history(history))
    except AppError:
        raise
    except Exception:
        raise AppError('invalid_excel', 'Повреждено содержимое Excel. Пересохраните файл .xlsx и повторите загрузку.', 422)
    finally:
        workbook.close()


def decode_file(raw, filename, kind):
    extension = Path(filename or '').suffix.lower()
    if extension == '.xlsx':
        return read_excel(raw)
    if extension == '.json':
        return read_json(raw, kind)
    if extension == '.csv' and kind == 'history':
        return None, raw
    raise AppError('unsupported_format', 'Поддерживаются Excel (.xlsx), JSON и CSV истории. Старый .xls пересохраните как .xlsx.', 422)


def merge_uploads(parts):
    employees = history = None
    for employee_part, history_part in parts:
        if employee_part is not None:
            if employees is not None:
                raise AppError('duplicate_source', 'Профили переданы в двух файлах. Оставьте один источник профилей.', 422)
            employees = employee_part
        if history_part is not None:
            if history is not None:
                raise AppError('duplicate_source', 'История передана в двух файлах. Оставьте один источник истории.', 422)
            history = history_part
    return employees, history
